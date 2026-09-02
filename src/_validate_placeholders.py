"""Throwaway helper: cross-check every layout's actual placeholders against
build_presentation.py's PAGE_FIELD_SPECS. Reports duplicate placeholder
names (breaks fill_slide's name->shape dict), missing required/optional
fields, and placeholders present but not in the spec. Read-only, changes
nothing. Not a project deliverable.

Usage: python src/_validate_placeholders.py <template.pptx> <out.txt>
"""

import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_presentation import PAGE_FIELD_SPECS  # noqa: E402
from pptx import Presentation  # noqa: E402


def main():
    template_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    prs = Presentation(str(template_path))

    lines = []
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            lines.append(f"=== {layout.name} ===")
            spec = PAGE_FIELD_SPECS.get(layout.name, {"required": [], "optional": []})
            expected = spec["required"] + spec["optional"]

            by_name = defaultdict(list)
            for ph in layout.placeholders:
                by_name[ph.name].append(ph.placeholder_format.idx)

            dupes = {name: idxs for name, idxs in by_name.items() if len(idxs) > 1}
            if dupes:
                for name, idxs in dupes.items():
                    lines.append(f"  DUPLICATE NAME: {name!r} used by idx {idxs} -- only one will be fillable")

            present = set(by_name.keys())
            missing_required = [f for f in spec["required"] if f not in present]
            missing_optional = [f for f in spec["optional"] if f not in present]
            unexpected = sorted(present - set(expected))

            if missing_required:
                lines.append(f"  MISSING REQUIRED: {missing_required}")
            if missing_optional:
                lines.append(f"  missing optional: {missing_optional}")
            if unexpected:
                lines.append(f"  present but not in PAGE_FIELD_SPECS: {unexpected}")
            if not (dupes or missing_required or missing_optional or unexpected):
                lines.append("  OK - matches spec exactly, no duplicates")

    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
