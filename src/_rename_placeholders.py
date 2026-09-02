"""Throwaway helper: rename generic PowerPoint placeholder shapes (by
layout name + placeholder idx) to the project's field-catalog names, in
place. Not a project deliverable.

Usage: python src/_rename_placeholders.py <template.pptx> <layout_name> <idx>=<new_name> [<idx>=<new_name> ...]
"""

import sys
from pathlib import Path

from pptx import Presentation


def main():
    template_path = Path(sys.argv[1])
    layout_name = sys.argv[2]
    renames = {}
    for arg in sys.argv[3:]:
        idx_str, new_name = arg.split("=", 1)
        renames[int(idx_str)] = new_name

    prs = Presentation(str(template_path))
    layout = None
    for master in prs.slide_masters:
        for l in master.slide_layouts:
            if l.name == layout_name:
                layout = l
    if layout is None:
        raise SystemExit(f"layout not found: {layout_name}")

    for shape in layout.placeholders:
        idx = shape.placeholder_format.idx
        if idx in renames:
            old = shape.name
            shape.name = renames[idx]
            print(f"idx={idx}: {old!r} -> {renames[idx]!r}")
            del renames[idx]

    if renames:
        print("WARNING: idx not found on this layout:", renames)

    prs.save(str(template_path))
    print("saved", template_path)


if __name__ == "__main__":
    main()
