"""Throwaway helper: dump every shape's name/position/size/placeholder-info
per layout in a given .pptx, to a UTF-8 text file (avoids Windows console
cp1251 crashes on German/Cyrillic text). Not a project deliverable.

Usage: python src/_dump_layout_shapes.py <template.pptx> <out.txt>
"""

import sys
from pathlib import Path

from pptx import Presentation

EMU_PER_CM = 360000


def main():
    template_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])
    prs = Presentation(str(template_path))

    lines = []
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            lines.append(f"=== {layout.name} ===")
            for shape in layout.shapes:
                l = shape.left / EMU_PER_CM if shape.left is not None else None
                t = shape.top / EMU_PER_CM if shape.top is not None else None
                w = shape.width / EMU_PER_CM if shape.width is not None else None
                h = shape.height / EMU_PER_CM if shape.height is not None else None
                ph_info = ""
                if shape.is_placeholder:
                    pf = shape.placeholder_format
                    ph_info = f" [PLACEHOLDER idx={pf.idx} type={pf.type}]"
                txt = ""
                if shape.has_text_frame and shape.text_frame.text.strip():
                    txt = " text=" + repr(shape.text_frame.text.strip()[:60])
                pos = f"pos=({l:.2f},{t:.2f}) size=({w:.2f}x{h:.2f})" if l is not None else "pos=(?)"
                lines.append(f"  {shape.shape_type} name={shape.name!r} {pos}{ph_info}{txt}")

    out_path.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
