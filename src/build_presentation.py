"""WP8 — turn a presentation JSON file into a finished PowerPoint file.

Usage:
    python src/build_presentation.py --input data/example_01.json \\
        --template templates/master_v01.pptx --output output/

Aspect-ratio decision (per the task brief's WP8 note that this needs writing
down): picture placeholders in `master_v01.pptx` were sized in WP6 to the
frame ratios recorded in docs/04-fields.md, but a real supplied photo won't
match those ratios exactly. `Placeholder.insert_picture()` crops to fill,
centred — fine for a close ratio, but a badly-mismatched one crops away part
of the motif. This script center-crops every image with Pillow to the exact
target placeholder ratio *before* calling insert_picture(), so the later
crop-to-fill is always a no-op. Chosen over "validate ratio and reject
mismatches" because the generator's whole point is to run unattended on
supplied photos without a human re-cropping them first.
"""

import argparse
import copy
import json
import sys
from pathlib import Path

from jsonschema import Draft202012Validator
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = PROJECT_ROOT / "schema" / "presentation.schema.json"

# PAGE_06_LOCAL_MOTIFS's headline/subheadline/stats-callout are computed from
# this page's position among the deck's PAGE_06_LOCAL_MOTIFS occurrences, not
# authored in the input JSON (docs/04-fields.md "Fields not included here",
# wording confirmed by the designer, docs/open-questions.md #22/#23).
# Position 4+ gets neither headline/subheadline nor the stats callout — only
# the theme-world-name caption (already an authored field, handled normally).
# Known scope gap: this doesn't model the row-limit-continuation mechanism
# (same theme overflowing onto a 2nd slide, e.g. Freiburg's 5-motif table) —
# the schema has no way to say "this PAGE_06 entry continues the previous
# one", so every entry is treated as a new theme at its sequence position.
MOTIF_POSITION_CONTENT = {
    1: {
        "TXT_MOTIF_HEADLINE": "Das nehmen Ihre Besucher mit — personalisiert, sofort, teilbar.",
        "TXT_MOTIF_SUBHEADLINE": (
            "Jedes Motiv wird individuell auf Ihren Weihnachtsmarkt abgestimmt — "
            "Ihre Besucher werden Teil Ihrer Erlebniswelt. Das teilen sie."
        ),
        "TXT_MOTIF_STATS_CALLOUT": (
            "70 % der Besucher teilen ihr Motiv aktiv auf Facebook, Instagram "
            "oder TikTok — mit Ihrem Markt als Kontext."
        ),
    },
    2: {
        "TXT_MOTIF_HEADLINE": "Ihre Charaktere. Ihre Geschichte.",
        "TXT_MOTIF_SUBHEADLINE": (
            "Gemeinsam mit Ihnen entwickeln wir die Charaktere, die perfekt zu Ihrem Markt passen."
        ),
        "TXT_MOTIF_STATS_CALLOUT": "",
    },
    3: {
        "TXT_MOTIF_HEADLINE": "Ein Weihnachtsmarkt. Mehrere Erlebniswelten.",
        "TXT_MOTIF_SUBHEADLINE": (
            "Jedes Motiv erzählt eine eigene Geschichte — perfekt abgestimmt "
            "auf Ihre Veranstaltung und Ihre Stadt."
        ),
        "TXT_MOTIF_STATS_CALLOUT": "",
    },
}

# Which fields each page_type can carry, split into required/optional, kept
# in sync with schema/presentation.schema.json's fields_PAGE_* defs. The
# schema is the source of truth for validation; this list drives *filling*
# behaviour (required-but-missing -> <<MISSING>> marker, optional-but-absent
# -> remove the unused placeholder so it doesn't show up empty).
PAGE_FIELD_SPECS = {
    "PAGE_01_TITLE": {
        "required": ["TXT_TITLE_SUBLINE", "IMG_THEME_WORLD_PHOTO"],
        "optional": [],
    },
    "PAGE_02_SERVICE": {
        "required": ["IMG_THEME_WORLD_PHOTO"],
        "optional": [],
    },
    "PAGE_03_THEME_SHOWCASE": {
        "required": ["TXT_THEME_WORLD_NAME_1", "IMG_SHOWCASE_BACKGROUND"],
        "optional": [
            "IMG_THEME_WORLD_PHOTO_1", "TXT_THEME_TAGLINE_1",
            "TXT_THEME_WORLD_NAME_2", "IMG_THEME_WORLD_PHOTO_2", "TXT_THEME_TAGLINE_2",
        ],
    },
    "PAGE_04_REFERENCES": {"required": [], "optional": []},
    "PAGE_05_USER_FLOW": {
        "required": [
            "TXT_KIOSK_BACKGROUND_LABEL",
            "IMG_KIOSK_SCREENSHOT",
            "IMG_USER_FLOW_CARD_1",
            "IMG_USER_FLOW_CARD_2",
        ],
        "optional": [],
    },
    "PAGE_06_LOCAL_MOTIFS": {
        "required": [
            "TXT_THEME_WORLD_NAME",
            "TXT_MOTIF_1_NAME",
            "IMG_MOTIF_1_PHOTO",
            "IMG_MOTIF_1_BACKGROUND",
            "IMG_MOTIF_1_MASK",
        ],
        "optional": [
            "TXT_MOTIF_2_NAME", "IMG_MOTIF_2_PHOTO", "IMG_MOTIF_2_BACKGROUND", "IMG_MOTIF_2_MASK",
            "TXT_MOTIF_3_NAME", "IMG_MOTIF_3_PHOTO", "IMG_MOTIF_3_BACKGROUND", "IMG_MOTIF_3_MASK",
        ],
    },
    "PAGE_07_BESTSELLERS": {"required": [], "optional": []},
    "PAGE_08_TRANSITION": {
        "required": ["TXT_THEME_WORLD_NAME_1", "IMG_OUTPUT_CARD_1", "TXT_THEME_WORLD_NAME_2", "IMG_OUTPUT_CARD_2"],
        "optional": [],
    },
    "PAGE_09_SOCIAL_REACH": {
        "required": ["IMG_INSTAGRAM_MOCKUP"],
        "optional": [],
    },
    "PAGE_10_CONTACT": {
        "required": ["IMG_LANDMARK_PHOTO"],
        "optional": [],
    },
}


def _make_console_utf8_safe():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def parse_args():
    p = argparse.ArgumentParser(description="Build a presentation .pptx from a JSON data file.")
    p.add_argument("--input", required=True, type=Path)
    p.add_argument("--template", required=True, type=Path)
    p.add_argument("--output", required=True, type=Path)
    return p.parse_args()


def load_and_validate(input_path: Path) -> dict:
    with open(input_path, encoding="utf-8") as f:
        data = json.load(f)
    with open(SCHEMA_PATH, encoding="utf-8") as f:
        schema = json.load(f)
    validator = Draft202012Validator(schema)
    errors = sorted(validator.iter_errors(data), key=lambda e: list(e.path))
    if errors:
        print(f"ERROR: {input_path} failed schema validation ({len(errors)} error(s)):", file=sys.stderr)
        for e in errors:
            location = "/".join(str(p) for p in e.path) or "(root)"
            print(f"  - {location}: {e.message}", file=sys.stderr)
        sys.exit(1)
    return data


def get_layout(prs: Presentation, page_type: str):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == page_type:
                return layout
    return None


def crop_to_ratio(image_path: Path, target_ratio: float, tmp_dir: Path) -> Path:
    """Center-crop image_path to target_ratio (width/height), return the path
    to use for insertion (the original, unmodified, if it already matches)."""
    im = Image.open(image_path)
    w, h = im.size
    current_ratio = w / h
    if abs(current_ratio - target_ratio) < 0.01:
        return image_path
    if current_ratio > target_ratio:
        new_w = round(h * target_ratio)
        x0 = (w - new_w) // 2
        box = (x0, 0, x0 + new_w, h)
    else:
        new_h = round(w / target_ratio)
        y0 = (h - new_h) // 2
        box = (0, y0, w, y0 + new_h)
    cropped = im.crop(box)
    tmp_dir.mkdir(parents=True, exist_ok=True)
    out_path = tmp_dir / f"{image_path.stem}_cropped{image_path.suffix}"
    cropped.save(out_path)
    return out_path


SLIDE_BACKGROUND_RGB = RGBColor(0xF2, 0xEA, 0xE0)


def remove_shape(shape):
    shape._element.getparent().remove(shape._element)


def _overlaps(a_left, a_top, a_width, a_height, b_left, b_top, b_width, b_height):
    a_right, a_bottom = a_left + a_width, a_top + a_height
    b_right, b_bottom = b_left + b_width, b_top + b_height
    return a_left < b_right and b_left < a_right and a_top < b_bottom and b_top < a_bottom


def crop_patch_from_full_bleed(source_path, frame_width, frame_height, patch_left, patch_top,
                                patch_width, patch_height, tmp_dir):
    """Crop the sub-region of a full-bleed image (already sized to exactly
    fill an EMU frame of frame_width x frame_height) that corresponds to a
    patch_left/top/width/height EMU box within that same frame, so the patch
    can be drawn back over the frame with no visible seam."""
    im = Image.open(source_path)
    img_w, img_h = im.size
    x0 = max(0, round(patch_left / frame_width * img_w))
    y0 = max(0, round(patch_top / frame_height * img_h))
    x1 = min(img_w, round((patch_left + patch_width) / frame_width * img_w))
    y1 = min(img_h, round((patch_top + patch_height) / frame_height * img_h))
    patch = im.crop((x0, y0, x1, y1))
    tmp_dir.mkdir(parents=True, exist_ok=True)
    out_path = tmp_dir / f"{Path(source_path).stem}_patch_{patch_left}_{patch_top}.png"
    patch.save(out_path)
    return out_path


def mask_orphaned_fixed_shapes(slide, layout, removed_bbox, background_patch=None, tmp_dir=None):
    """Cover any of the layout's fixed (non-placeholder) shapes that overlap
    removed_bbox, hiding them behind whatever the slide's real background is.

    A layout can have a fixed decorative shape (e.g. a tinted caption plate)
    positioned specifically behind one optional placeholder slot — normally
    hidden by that slot's own opaque content. Once the slot's placeholder is
    removed for being unused (see fill_slide), that decoration has nothing
    left covering it and floats into view via ordinary layout inheritance.
    Found on PAGE_03_THEME_SHOWCASE's unused 2nd theme slot. Masking with a
    plain rectangle is simpler and more robust than trying to hide the
    layout shape itself (layout shapes are shared across every slide using
    that layout, so a single slide can't delete one).

    A flat SLIDE_BACKGROUND_RGB rectangle only blends in on layouts whose
    real background is that plain cream fill. On a layout whose background
    is itself a full-bleed photo (PAGE_03_THEME_SHOWCASE, once its
    IMG_SHOWCASE_BACKGROUND field is filled), a flat-colour cover shows up
    as an obviously mismatched box instead of blending in — pass
    background_patch (the (image_path, left, top, width, height) of that
    full-bleed photo, as already inserted) and this crops the matching
    region of the same photo to patch over the orphan instead.

    Returns the set of layout shape_ids that were covered, so a caller that
    later runs reassert_fixed_layout_shapes() on the same slide (layouts
    that need both, e.g. PAGE_03_THEME_SHOWCASE) can tell it to skip
    re-copying these particular shapes fresh from the layout — otherwise
    that blind re-copy would paste the orphan right back on top of the
    patch this function just added.
    """
    left, top, width, height = removed_bbox
    masked_shape_ids = set()
    for shape in layout.shapes:
        if shape.is_placeholder:
            continue
        if not _overlaps(left, top, width, height, shape.left, shape.top, shape.width, shape.height):
            continue
        masked_shape_ids.add(shape.shape_id)
        if background_patch is not None:
            source_path, frame_left, frame_top, frame_width, frame_height = background_patch
            patch_path = crop_patch_from_full_bleed(
                source_path, frame_width, frame_height,
                shape.left - frame_left, shape.top - frame_top, shape.width, shape.height, tmp_dir,
            )
            slide.shapes.add_picture(str(patch_path), shape.left, shape.top, shape.width, shape.height)
        else:
            cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, shape.left, shape.top, shape.width, shape.height)
            cover.fill.solid()
            cover.fill.fore_color.rgb = SLIDE_BACKGROUND_RGB
            cover.line.fill.background()
            cover.shadow.inherit = False
    return masked_shape_ids


# Layouts where a fixed (non-placeholder) shape needs to render in front of
# a full-bleed photo placeholder — see reassert_fixed_layout_shapes().
# PAGE_01_TITLE: its headline is a plain fixed TextBox with no placeholder
# equivalent, so it's invisible behind any inserted photo unless re-asserted.
# PAGE_03_THEME_SHOWCASE: added once its redesigned layout gave
# IMG_SHOWCASE_BACKGROUND a full-bleed frame too (it didn't have one before),
# which hid its name/tagline plates the exact same way. Unlike PAGE_01, this
# layout also has an *unused-slot* masking step (mask_orphaned_fixed_shapes,
# for the common single-theme case) that runs first and covers slot 2's
# now-orphaned plate — reassert must be told to skip re-copying exactly
# those already-covered shapes (via masked_shape_ids), or its blind copy
# would paste the orphan right back on top of that cover.
LAYOUTS_NEEDING_FIXED_SHAPE_REASSERT = {"PAGE_01_TITLE", "PAGE_02_SERVICE", "PAGE_03_THEME_SHOWCASE"}

# EMU tolerance for treating a fixed shape's bbox as "the same slot" as a
# placeholder's, in reassert_fixed_layout_shapes() below.
PLATE_MATCH_TOLERANCE = 20000


def reassert_fixed_layout_shapes(slide, layout, skip_shape_ids=frozenset()):
    """Copy the layout's own non-placeholder ("fixed content") shapes onto
    the slide.

    PowerPoint renders layout-inherited fixed shapes behind ALL of a slide's
    own placeholder content, regardless of their relative order in the
    layout's XML. That's invisible on most layouts (nothing overlaps them),
    but on a layout with a full-bleed (or near-full-bleed) photo placeholder,
    that photo completely covers any fixed shape behind it — the headline
    text on PAGE_01_TITLE, the theme-name caption plate on PAGE_03, the
    caption plate on PAGE_02_SERVICE's near-full-bleed circular photo — found
    the hard way while sanity-checking real generated decks. Explicit copies
    here render in front of the inherited (now-hidden) originals, which is
    harmless — same content, same position, just no longer hidden.

    Not every fixed shape wants the same *stacking* position once re-copied,
    though. A shape with no placeholder equivalent (e.g. PAGE_01's headline
    TextBox) has nothing else to layer against, so it simply goes frontmost.
    But a fixed shape that's really a tinted *backdrop plate* for one of the
    slide's own text placeholders — same position and size, added later so
    the caption reads better over a busy photo — must render BEHIND that
    placeholder's text, not in front of it (else the plate's fill visually
    covers the very text it's supposed to set off). So: for each fixed shape,
    look for a same-bbox text placeholder already on the slide; if found,
    insert the copy immediately before that placeholder's own element
    (behind it, but still in front of the photo, which sits earlier in the
    tree); otherwise append at the end (frontmost) as before.

    skip_shape_ids excludes specific layout shapes (by shape_id) from being
    re-copied — needed for a layout like PAGE_03_THEME_SHOWCASE where an
    earlier masking step has already covered an orphaned shape on this
    slide; blindly re-copying it here would undo that cover.

    A plain XML deepcopy is enough for text/shape content, but a fixed
    *picture* shape's `<a:blip r:embed="rIdN">` points at a relationship
    that only exists in the layout part's own .rels — copying the element
    as-is into the slide leaves that rId dangling there (broken image, or
    by chance colliding with an unrelated rId already used on the slide).
    So for every blip found in the copied element, re-resolve the source
    image part from the layout part and add a fresh relationship on the
    slide part, pointing the copy's r:embed at the new, valid rId.
    """
    # Snapshot text-placeholder bboxes up front: reading .left/.top/.width/
    # .height on a PICTURE placeholder after insert_picture() has already
    # run on it (fill_slide runs before this) raises inside python-pptx, and
    # we only ever need to backdrop-match text placeholders anyway.
    text_placeholders = [
        ph for ph in slide.placeholders
        if ph.placeholder_format.type != PP_PLACEHOLDER.PICTURE
    ]
    text_placeholder_bboxes = [
        (ph, ph.left, ph.top, ph.width, ph.height) for ph in text_placeholders
    ]

    for shape in layout.shapes:
        if shape.is_placeholder or shape.shape_id in skip_shape_ids:
            continue
        element = copy.deepcopy(shape._element)
        for blip in element.iter(qn("a:blip")):
            old_rid = blip.get(qn("r:embed"))
            if not old_rid:
                continue
            image_part = layout.part.related_part(old_rid)
            new_rid = slide.part.relate_to(image_part, RT.IMAGE)
            blip.set(qn("r:embed"), new_rid)

        target = None
        for ph, ph_left, ph_top, ph_width, ph_height in text_placeholder_bboxes:
            if (
                abs(ph_left - shape.left) <= PLATE_MATCH_TOLERANCE
                and abs(ph_top - shape.top) <= PLATE_MATCH_TOLERANCE
                and abs(ph_width - shape.width) <= PLATE_MATCH_TOLERANCE
                and abs(ph_height - shape.height) <= PLATE_MATCH_TOLERANCE
            ):
                target = ph
                break

        if target is not None:
            target._element.addprevious(element)
        else:
            slide.shapes._spTree.append(element)


def fill_slide(slide, page_type: str, fields: dict, base_dir: Path, tmp_dir: Path, log,
               slide_width: int, slide_height: int):
    name_by_idx = {ph.placeholder_format.idx: ph.name for ph in slide.slide_layout.placeholders}
    placeholders = {name_by_idx.get(ph.placeholder_format.idx): ph for ph in slide.placeholders}
    spec = PAGE_FIELD_SPECS.get(page_type, {"required": [], "optional": []})

    # Tracks the (image_path, left, top, width, height) of an image field
    # that turned out to be full-bleed (covers the whole slide), so any
    # orphaned decorative shape masked below can be patched with the
    # matching crop of that same photo instead of a flat colour rectangle
    # — see mask_orphaned_fixed_shapes(). None if no full-bleed photo was
    # filled on this slide (e.g. it's still <<MISSING>> or wasn't provided).
    full_bleed_patch = None
    bleed_tolerance = 10000  # EMU (~0.01in) of slack for rounding
    masked_shape_ids = set()

    for field_name in spec["required"] + spec["optional"]:
        is_required = field_name in spec["required"]
        placeholder = placeholders.get(field_name)
        if placeholder is None:
            log(f"    ! template has no placeholder named {field_name} (schema/template drift)")
            continue

        value = fields.get(field_name)
        is_image = field_name.startswith("IMG_")

        if value is not None:
            if is_image:
                image_path = base_dir / value
                if not image_path.exists():
                    log(f"    WARNING: {field_name} -> image file not found: {image_path}")
                    continue
                ph_left, ph_top, ph_width, ph_height = (
                    placeholder.left, placeholder.top, placeholder.width, placeholder.height,
                )
                ratio = ph_width / ph_height
                final_path = crop_to_ratio(image_path, ratio, tmp_dir)
                placeholder.insert_picture(str(final_path))
                log(f"    filled {field_name} <- {value}")
                if (
                    abs(ph_left) <= bleed_tolerance
                    and abs(ph_top) <= bleed_tolerance
                    and abs(ph_width - slide_width) <= bleed_tolerance
                    and abs(ph_height - slide_height) <= bleed_tolerance
                ):
                    full_bleed_patch = (final_path, ph_left, ph_top, ph_width, ph_height)
            else:
                placeholder.text_frame.text = value
                log(f"    filled {field_name} = {value!r}")
        else:
            if is_required:
                log(f"    WARNING: required field {field_name} missing")
                if not is_image:
                    placeholder.text_frame.text = f"<<MISSING: {field_name}>>"
            else:
                log(f"    optional field {field_name} not provided, removing unused placeholder")
                bbox = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
                remove_shape(placeholder)
                masked_shape_ids |= mask_orphaned_fixed_shapes(
                    slide, slide.slide_layout, bbox, full_bleed_patch, tmp_dir,
                )

    return full_bleed_patch, masked_shape_ids


def fill_motif_computed_fields(slide, position: int, log):
    name_by_idx = {ph.placeholder_format.idx: ph.name for ph in slide.slide_layout.placeholders}
    placeholders = {name_by_idx.get(ph.placeholder_format.idx): ph for ph in slide.placeholders}
    content = MOTIF_POSITION_CONTENT.get(position)

    for field_name in ("TXT_MOTIF_HEADLINE", "TXT_MOTIF_SUBHEADLINE", "TXT_MOTIF_STATS_CALLOUT"):
        placeholder = placeholders.get(field_name)
        if placeholder is None:
            continue
        text = content.get(field_name, "") if content else ""
        if text:
            placeholder.text_frame.text = text
            log(f"    computed {field_name} (world position {position}) = {text!r}")
        else:
            log(f"    computed {field_name} (world position {position}) -> empty, removing placeholder")
            bbox = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
            remove_shape(placeholder)
            mask_orphaned_fixed_shapes(slide, slide.slide_layout, bbox)


# PAGE_04_REFERENCES's client-list footer is computed, not authored: the
# fixed 8-name boilerplate list minus the current deck's own market
# (docs/03-elements.md's "Client-list footer" row, docs/open-questions.md
# #5). Split into the source decks' original 2-line grouping so the normal
# (nothing excluded) case reproduces that layout; a self-match just drops
# out of whichever line it's on.
REFERENCES_FOOTER_CLIENTS_LINE_1 = [
    "Striezelmarkt Dresden",
    "Salzburger Christkindlmarkt",
    "Berliner Weihnachtszeit",
    "Weihnachtsmarkt Wiesbaden",
]
REFERENCES_FOOTER_CLIENTS_LINE_2 = [
    "Kölner Dom",
    "Tower Bridge London",
    "Checkpoint Charlie Berlin",
    "Wiener Prater",
]


def compute_references_footer(city: str) -> str:
    def exclude_self(names):
        return [n for n in names if city.lower() not in n.lower()]

    line1 = "     ".join(exclude_self(REFERENCES_FOOTER_CLIENTS_LINE_1))
    line2 = "     ".join(exclude_self(REFERENCES_FOOTER_CLIENTS_LINE_2) + ["und viele mehr."])
    return f"{line1}\n{line2}"


def fill_references_footer(slide, city: str, log):
    name_by_idx = {ph.placeholder_format.idx: ph.name for ph in slide.slide_layout.placeholders}
    placeholders = {name_by_idx.get(ph.placeholder_format.idx): ph for ph in slide.placeholders}
    placeholder = placeholders.get("TXT_REFERENCES_FOOTER")
    if placeholder is None:
        log("    ! template has no placeholder named TXT_REFERENCES_FOOTER (schema/template drift)")
        return
    text = compute_references_footer(city)
    placeholder.text_frame.text = text
    log(f"    computed TXT_REFERENCES_FOOTER = {text!r}")


# PAGE_01_TITLE/PAGE_02_SERVICE's theme-world-name caption is likewise
# computed, not authored per-page: theme_worlds[0], shown only once the deck
# has more than 2 theme worlds (docs/open-questions.md #12, top-level
# theme_worlds' schema description). Below that threshold the placeholder is
# removed, same optional-field-removal treatment as fill_slide's.
LAYOUTS_WITH_COMPUTED_THEME_CAPTION = {"PAGE_01_TITLE", "PAGE_02_SERVICE"}


def fill_theme_world_name_caption(slide, layout, theme_worlds: list, log, background_patch=None, tmp_dir=None):
    name_by_idx = {ph.placeholder_format.idx: ph.name for ph in slide.slide_layout.placeholders}
    placeholders = {name_by_idx.get(ph.placeholder_format.idx): ph for ph in slide.placeholders}
    placeholder = placeholders.get("TXT_THEME_WORLD_NAME")
    if placeholder is None:
        return

    if len(theme_worlds) > 2:
        text = theme_worlds[0]
        placeholder.text_frame.text = text
        log(f"    computed TXT_THEME_WORLD_NAME = {text!r}")
    else:
        log("    computed TXT_THEME_WORLD_NAME -> deck has <=2 theme worlds, removing unused placeholder")
        bbox = (placeholder.left, placeholder.top, placeholder.width, placeholder.height)
        remove_shape(placeholder)
        mask_orphaned_fixed_shapes(slide, layout, bbox, background_patch, tmp_dir)


def next_output_path(output_dir: Path, project_id: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    n = 1
    while True:
        candidate = output_dir / f"{project_id}_v{n:02d}.pptx"
        if not candidate.exists():
            return candidate
        n += 1


def main():
    _make_console_utf8_safe()
    args = parse_args()

    def log(msg):
        print(msg)

    log(f"Loading and validating {args.input} against {SCHEMA_PATH} ...")
    data = load_and_validate(args.input)
    log("Validation passed.")

    base_dir = args.input.resolve().parent.parent  # data/ -> project root, so 'images/...' resolves
    tmp_dir = PROJECT_ROOT / "output" / "_build_tmp"

    log(f"Opening template {args.template} ...")
    prs = Presentation(str(args.template))

    motif_table_count = 0
    for i, page in enumerate(data["pages"], start=1):
        page_type = page["page_type"]
        fields = page.get("fields", {})
        log(f"Slide {i}: {page_type}")

        layout = get_layout(prs, page_type)
        if layout is None:
            log(f"  ERROR: template has no layout named {page_type}, skipping")
            continue

        slide = prs.slides.add_slide(layout)
        full_bleed_patch, masked_shape_ids = fill_slide(slide, page_type, fields, base_dir, tmp_dir, log,
                                                          prs.slide_width, prs.slide_height)
        if page_type in LAYOUTS_NEEDING_FIXED_SHAPE_REASSERT:
            reassert_fixed_layout_shapes(slide, layout, masked_shape_ids)

        if page_type == "PAGE_06_LOCAL_MOTIFS":
            motif_table_count += 1
            fill_motif_computed_fields(slide, motif_table_count, log)
        elif page_type == "PAGE_04_REFERENCES":
            fill_references_footer(slide, data["city"], log)
        elif page_type in LAYOUTS_WITH_COMPUTED_THEME_CAPTION:
            fill_theme_world_name_caption(slide, layout, data["theme_worlds"], log, full_bleed_patch, tmp_dir)

    out_path = next_output_path(args.output, data["project_id"])
    prs.save(str(out_path))
    log(f"Saved {out_path}")


if __name__ == "__main__":
    main()
